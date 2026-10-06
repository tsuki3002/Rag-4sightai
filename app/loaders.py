"""Turn source files into a list of (page/section, text) units so citations can point somewhere precise."""
import logging
from dataclasses import dataclass
from pathlib import Path

log = logging.getLogger(__name__)

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
SUPPORTED = {".pdf", ".md", ".txt", ".docx", ".pptx", ".html", ".htm"} | IMAGE_EXT
SCANNED_PAGE_CHARS = 30  # PDF pages with less extractable text than this get OCR'd
# reference-list sections add noise to search results without answering anything
BOILERPLATE_SECTIONS = {"references", "external links", "further reading", "see also", "notes", "bibliography",
                        "sources", "citations", "footnotes"}


@dataclass
class Unit:
    text: str
    page: int | None = None      # PDFs: page number; PPTX: slide number
    section: str | None = None   # Markdown/DOCX/HTML: nearest heading; PPTX: slide title


def load(path: Path) -> list[Unit]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _load_pdf(path)
    if ext == ".md":
        return _load_markdown(path.read_text(encoding="utf-8", errors="ignore"))
    if ext == ".txt":
        return [Unit(text=path.read_text(encoding="utf-8", errors="ignore"))]
    if ext == ".docx":
        return _load_docx(path)
    if ext == ".pptx":
        return _load_pptx(path)
    if ext in {".html", ".htm"}:
        return _load_html(path.read_text(encoding="utf-8", errors="ignore"))
    if ext in IMAGE_EXT:
        return _load_image(path)
    raise ValueError(f"Unsupported file type: {path}")


def _load_pdf(path: Path) -> list[Unit]:
    from pypdf import PdfReader

    from .vision import read_image

    units = []
    for i, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = page.extract_text() or ""
        if len(text.strip()) < SCANNED_PAGE_CHARS:
            # scanned page: the "text" is a picture, so OCR the page's images
            try:
                ocr = [read_image(img.image, f"{path.name} page {i}") for img in page.images]
                text = "\n".join(t for t in [text, *ocr] if t.strip())
            except Exception as e:
                log.warning("Couldn't OCR %s page %d: %s", path.name, i, e)
        units.append(Unit(text=text, page=i))
    return units


def _load_image(path: Path) -> list[Unit]:
    from PIL import Image

    from .vision import read_image

    with Image.open(path) as img:
        return [Unit(text=read_image(img, path.stem.replace("_", " ")))]


def _load_pptx(path: Path) -> list[Unit]:
    import io

    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from .vision import read_image

    units = []
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else None
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                parts.extend(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    parts.append(read_image(Image.open(io.BytesIO(shape.image.blob)), f"slide {n} figure"))
                except Exception as e:
                    log.warning("Couldn't read image on %s slide %d: %s", path.name, n, e)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
        units.append(Unit(text="\n".join(p for p in parts if p), page=n, section=title))
    return units


def _load_html(html: str) -> list[Unit]:
    import trafilatura

    text = trafilatura.extract(html, output_format="markdown", include_tables=True, include_comments=False)
    return _load_markdown(text or "")


def _load_markdown(text: str) -> list[Unit]:
    lines, in_fence = [], False
    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            in_fence = not in_fence
        is_heading = not in_fence and line.lstrip().startswith("#")
        lines.append((line, line.lstrip("# ").strip().rstrip("¶").strip() if is_heading else None))
    return _group_by_heading(lines)


def _load_docx(path: Path) -> list[Unit]:
    from docx import Document

    def is_heading(p) -> bool:
        return p.style is not None and p.style.name.lower().startswith("heading")

    return _group_by_heading(
        (p.text, p.text.strip() if is_heading(p) else None) for p in Document(str(path)).paragraphs
    )


def _group_by_heading(lines) -> list[Unit]:
    """`lines` yields (text, heading) where heading is set only on heading lines."""
    units, section, buf = [], None, []

    def flush():
        text = "\n".join(buf).strip()
        if text and (section or "").strip().lower() not in BOILERPLATE_SECTIONS:
            units.append(Unit(text=text, section=section))

    for line, heading in lines:
        if heading is not None:
            flush()
            buf = [line]
            section = heading
        else:
            buf.append(line)
    flush()
    return units
